"""
DOCX parsing and processing utilities.
"""
import re
from docx import Document
from typing import List


class DOCXParser:
    """Service for parsing and processing DOCX documents."""
    
    @staticmethod
    def extract_placeholders(doc: Document) -> List[str]:
        """Extract all placeholders from document including tables."""
        placeholders = []

        def iter_paragraphs(doc_or_cell):
            for para in doc_or_cell.paragraphs:
                yield para
            for table in doc_or_cell.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from iter_paragraphs(cell)

        for para in iter_paragraphs(doc):
            matches = re.findall(r"<(.*?)>", para.text)
            placeholders.extend(matches)

        return placeholders

    @staticmethod
    def get_context_type(paragraph) -> str:
        """Determine if paragraph is in table or section context."""
        return "table" if paragraph._element.getparent().tag.endswith("tc") else "section"
