"""
Template processing routes.
"""
from fastapi import APIRouter, Depends, HTTPException, BackgroundTasks
from sqlmodel import Session
from pydantic import BaseModel
from typing import Optional, List
import os
import uuid
import asyncio
from utils.database import get_session
from services.template_filler import TemplateFiller, extract_placeholders, retrieve_placeholder_content
from services.llm_service import LLMService
from utils.config import INPUT_FOLDER, GENERATED_FOLDER
from docx import Document
from api.websocket import broadcast_progress_update_sync

router = APIRouter(prefix="/api/template", tags=["template"])


class TemplateRequest(BaseModel):
    folder_name: str
    user_prompt: Optional[str] = ""
    process_flow: Optional[str] = ""
    selected_files: Optional[List[str]] = None  # List of specific files to process


class TemplateResponse(BaseModel):
    task_id: str
    message: str


# Store for tracking template processing tasks
template_tasks = {}


@router.post("/fill", response_model=TemplateResponse)
def start_template_filling(
    request: TemplateRequest, 
    background_tasks: BackgroundTasks,
    session: Session = Depends(get_session)
):
    """Start template filling process and return task ID."""
    folder_name = request.folder_name
    folder_path = os.path.join(INPUT_FOLDER, folder_name)

    print(f"[TEMPLATE FILL] Processing request for folder: {request.folder_name}")

    if not os.path.exists(folder_path) or not os.path.isdir(folder_path):
        raise HTTPException(status_code=404, detail=f"Folder '{folder_name}' not found in inputs")

    # Generate unique task ID
    task_id = str(uuid.uuid4())
    
    # Initialize task tracking
    template_tasks[task_id] = {
        "status": "processing",
        "files_done": 0,
        "files_total": 0,
        "generated_files": []
    }

    # Start background processing
    background_tasks.add_task(
        process_templates_background,
        task_id,
        folder_path,
        request.user_prompt,
        request.process_flow,
        session,
        request.selected_files
    )

    return TemplateResponse(
        task_id=task_id,
        message="Template processing started"
    )


def process_templates_background(
    task_id: str,
    folder_path: str,
    user_prompt: str,
    process_flow: str,
    session: Session,
    selected_files: List[str] = None
):
    """Background task to process templates."""
    print(f"[TEMPLATE PROCESSING] Starting task {task_id}")
    try:
        # Generate process flow description if provided
        process_flow_description = None
        if process_flow:
            llm_service = LLMService()
            process_flow_description = llm_service.generate_process_flow_description(process_flow)

        # Get .docx files to process
        all_docx_files = [f for f in os.listdir(folder_path) if f.endswith(".docx")]
        
        # Filter by selected files if provided
        if selected_files:
            # Check which selected files actually exist in the folder
            existing_selected = [f for f in selected_files if f in all_docx_files]
            missing_files = [f for f in selected_files if f not in all_docx_files]
            
            if existing_selected:
                docx_files = existing_selected
                print(f"[TEMPLATE PROCESSING] Processing {len(existing_selected)} selected files")
            else:
                docx_files = all_docx_files
                print(f"[TEMPLATE PROCESSING] No selected files found, processing all {len(all_docx_files)} files")
        else:
            docx_files = all_docx_files
            print(f"[TEMPLATE PROCESSING] Processing all {len(all_docx_files)} files")
            
        template_tasks[task_id]["files_total"] = len(docx_files)

        # Send initial call log
        try:
            log_data = {
                "type": "call_log",
                "service": "template_processor",
                "message": f"Starting to process {len(docx_files)} template files",
                "logType": "info"
            }
            broadcast_progress_update_sync(task_id, log_data)
        except Exception as ws_error:
            print(f"[TEMPLATE PROCESSING] WebSocket error (start): {ws_error}")

        # Process each template
        for file_name in docx_files:
            
            # Send call log for file start
            try:
                log_data = {
                    "type": "call_log",
                    "service": "template_processor",
                    "message": f"Processing template: {file_name}",
                    "logType": "info"
                }
                broadcast_progress_update_sync(task_id, log_data)
            except Exception as ws_error:
                print(f"[TEMPLATE PROCESSING] WebSocket error (start): {ws_error}")
            
            try:
                file_path = os.path.join(folder_path, file_name)
                doc = Document(file_path)
                placeholders = extract_placeholders(doc)
                filler = TemplateFiller(task_id=task_id)

                # Create retrieve function with context
                def retrieve_fn(ph, context_type):
                    return retrieve_placeholder_content(
                        ph,
                        context_type,
                        session,
                        user_prompt=user_prompt,
                        process_flow=process_flow_description,
                        task_id=task_id
                    )

                # Fill placeholders
                filled_doc = filler.fill_placeholders(doc, placeholders, retrieve_fn)

                # Save filled document
                output_filename = f"filled_{file_name}"
                output_path = os.path.join(GENERATED_FOLDER, output_filename)
                filled_doc.save(output_path)

                # Update task progress
                template_tasks[task_id]["files_done"] += 1
                # URL encode the filename for proper download URL
                from urllib.parse import quote
                encoded_filename = quote(output_filename)
                template_tasks[task_id]["generated_files"].append({
                    "fileName": output_filename,
                    "status": "done",
                    "downloadUrl": f"/api/template/download/{encoded_filename}"
                })
                
                # Send file completion update (legacy support)
                file_update = {
                    "fileName": output_filename,
                    "status": "done",
                    "downloadUrl": f"/api/template/download/{encoded_filename}",
                    "filesDone": template_tasks[task_id]["files_done"],
                    "filesTotal": template_tasks[task_id]["files_total"]
                }
                
                # Send call log for file completion
                try:
                    log_data = {
                        "type": "call_log",
                        "service": "template_processor",
                        "message": f"Completed: {output_filename}",
                        "logType": "success"
                    }
                    broadcast_progress_update_sync(task_id, log_data)
                    
                    # Send file update for download functionality
                    broadcast_progress_update_sync(task_id, file_update)
                except Exception as ws_error:
                    print(f"[TEMPLATE PROCESSING] WebSocket error: {ws_error}")

            except Exception as e:
                print(f"Error processing {file_name}: {e}")
                template_tasks[task_id]["generated_files"].append({
                    "fileName": file_name,
                    "status": "error",
                    "downloadUrl": None
                })

        # Send final completion message
        try:
            final_message = f"Template processing completed successfully! Generated {len(template_tasks[task_id]['generated_files'])} files."
            final_response = {
                "type": "final_response",
                "message": final_message
            }
            broadcast_progress_update_sync(task_id, final_response)
        except Exception as ws_error:
            print(f"[TEMPLATE PROCESSING] WebSocket error (final): {ws_error}")

        # Mark task as completed
        template_tasks[task_id]["status"] = "completed"
        print(f"[TEMPLATE PROCESSING] Task {task_id} completed successfully")

    except Exception as e:
        template_tasks[task_id]["status"] = "error"
        print(f"[TEMPLATE PROCESSING] Background task error: {e}")
        import traceback
        traceback.print_exc()


@router.get("/progress/{task_id}")
def get_template_progress(task_id: str):
    """Get progress of template processing task."""
    if task_id not in template_tasks:
        raise HTTPException(status_code=404, detail="Task not found")
    
    task = template_tasks[task_id]
    return {
        "task_id": task_id,
        "status": task["status"],
        "files_done": task["files_done"],
        "files_total": task["files_total"],
        "generated_files": task["generated_files"]
    }


@router.get("/download/{filename:path}")
def download_generated_file(filename: str):
    """Download a generated template file."""
    from urllib.parse import unquote
    
    # Decode URL-encoded filename
    decoded_filename = unquote(filename)
    
    # Try multiple possible paths
    possible_paths = [
        os.path.join(GENERATED_FOLDER, decoded_filename),  # Local path
        os.path.join("/app/generated", decoded_filename),   # Docker path
        os.path.join("./generated", decoded_filename),      # Relative path
        os.path.join("../generated", decoded_filename),     # Parent relative path
        os.path.join("/app", "generated", decoded_filename), # Alternative Docker path
    ]
    
    print(f"[DOWNLOAD] Requested filename: {filename}")
    print(f"[DOWNLOAD] Decoded filename: {decoded_filename}")
    
    # Check each possible path
    file_path = None
    for path in possible_paths:
        print(f"[DOWNLOAD] Checking path: {path} - Exists: {os.path.exists(path)}")
        if os.path.exists(path):
            file_path = path
            break
    
    # List all files in possible generated folders for debugging
    for folder_path in [GENERATED_FOLDER, "/app/generated", "./generated", "../generated"]:
        if os.path.exists(folder_path):
            try:
                generated_files = os.listdir(folder_path)
                print(f"[DOWNLOAD] Available files in {folder_path}: {generated_files}")
            except Exception as e:
                print(f"[DOWNLOAD] Error listing {folder_path}: {e}")
    
    if not file_path:
        available_paths = []
        for folder_path in [GENERATED_FOLDER, "/app/generated", "./generated", "../generated"]:
            if os.path.exists(folder_path):
                available_paths.append(folder_path)
        raise HTTPException(
            status_code=404, 
            detail=f"File not found: {decoded_filename}. Searched paths: {possible_paths}. Available folders: {available_paths}"
        )
    
    from fastapi.responses import FileResponse
    print(f"[DOWNLOAD] Serving file from: {file_path}")
    return FileResponse(
        path=file_path,
        filename=decoded_filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


@router.get("/list-templates/{folder_name}")
def list_template_files(folder_name: str):
    """List all available template files in a package."""
    folder_path = os.path.join(INPUT_FOLDER, folder_name)
    
    if not os.path.exists(folder_path) or not os.path.isdir(folder_path):
        raise HTTPException(status_code=404, detail=f"Folder '{folder_name}' not found in inputs")
    
    try:
        all_files = os.listdir(folder_path)
        docx_files = [f for f in all_files if f.endswith(".docx")]
        
        file_info = []
        for file in docx_files:
            file_path = os.path.join(folder_path, file)
            if os.path.isfile(file_path):
                stat = os.stat(file_path)
                file_info.append({
                    "filename": file,
                    "size_bytes": stat.st_size,
                    "size_mb": round(stat.st_size / (1024 * 1024), 2),
                    "modified": stat.st_mtime
                })
        
        return {
            "folder_name": folder_name,
            "message": f"Found {len(file_info)} template files",
            "files": file_info
        }
    except Exception as e:
        return {"error": f"Error listing template files: {str(e)}"}


@router.get("/test-download")
def test_download():
    """Test endpoint to check download functionality."""
    import glob
    
    # Check all possible generated folders
    possible_folders = [
        GENERATED_FOLDER,
        "/app/generated",
        "./generated", 
        "../generated",
        "/app/generated"
    ]
    
    result = {}
    for folder in possible_folders:
        if os.path.exists(folder):
            try:
                files = os.listdir(folder)
                result[folder] = {
                    "exists": True,
                    "files": files,
                    "count": len(files)
                }
            except Exception as e:
                result[folder] = {
                    "exists": True,
                    "error": str(e)
                }
        else:
            result[folder] = {"exists": False}
    
    return {
        "message": "Download test results",
        "folders": result,
        "generated_folder_config": GENERATED_FOLDER
    }


@router.get("/list-generated")
def list_generated_files():
    """List all generated files in the generated folder."""
    if not os.path.exists(GENERATED_FOLDER):
        return {"message": "Generated folder does not exist", "files": []}
    
    try:
        files = os.listdir(GENERATED_FOLDER)
        file_info = []
        for file in files:
            file_path = os.path.join(GENERATED_FOLDER, file)
            if os.path.isfile(file_path):
                stat = os.stat(file_path)
                file_info.append({
                    "filename": file,
                    "size_bytes": stat.st_size,
                    "size_mb": round(stat.st_size / (1024 * 1024), 2),
                    "download_url": f"/api/template/download/{file}"
                })
        
        return {
            "message": f"Found {len(file_info)} generated files",
            "files": file_info
        }
    except Exception as e:
        return {"error": f"Error listing files: {str(e)}"}


@router.post("/test-fill")
def test_template_fill(
    request: TemplateRequest, 
    session: Session = Depends(get_session)
):
    """Test template filling without background tasks."""
    folder_name = request.folder_name
    folder_path = os.path.join(INPUT_FOLDER, folder_name)

    if not os.path.exists(folder_path) or not os.path.isdir(folder_path):
        raise HTTPException(status_code=404, detail=f"Folder '{folder_name}' not found in inputs")

    try:
        # Get all .docx files in folder
        docx_files = [f for f in os.listdir(folder_path) if f.endswith(".docx")]
        
        if not docx_files:
            return {"message": "No .docx files found in the folder", "files": []}

        # Process first file as test
        file_name = docx_files[0]
        file_path = os.path.join(folder_path, file_name)
        doc = Document(file_path)
        placeholders = extract_placeholders(doc)
        
        # Generate process flow description if provided
        process_flow_description = None
        if request.process_flow:
            llm_service = LLMService()
            process_flow_description = llm_service.generate_process_flow_description(request.process_flow)

        # Create retrieve function with context
        def retrieve_fn(ph, context_type):
            return retrieve_placeholder_content(
                ph,
                context_type,
                session,
                user_prompt=request.user_prompt,
                process_flow=process_flow_description
            )

        # Fill placeholders
        filler = TemplateFiller()
        filled_doc = filler.fill_placeholders(doc, placeholders, retrieve_fn)

        # Save filled document
        output_filename = f"test_filled_{file_name}"
        output_path = os.path.join(GENERATED_FOLDER, output_filename)
        filled_doc.save(output_path)

        # URL encode the filename for proper download URL
        from urllib.parse import quote
        encoded_filename = quote(output_filename)
        
        return {
            "message": f"Successfully processed {file_name}",
            "output_file": output_filename,
            "placeholders_found": len(placeholders),
            "download_url": f"/api/template/download/{encoded_filename}"
        }

    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Error processing template: {str(e)}")
